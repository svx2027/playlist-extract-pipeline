#!/usr/bin/env python3
"""Build content_brain_vault.xlsx: one row per video with an embedded thumbnail,
an AI-written summary, and editable columns (with dropdowns) for the review loop.

This script is the CONSUMER half of the content-brain layer described in
docs/EXTRACTION_PIPELINE_PLAYBOOK.md, section 7. It expects two inputs that this
repo does not generate for you (the summarization fan-out is meant to run in
whatever LLM tooling you already have -- see the README's "Building the vault"
section for the exact shapes):

  out/brain_jobs.json          a JSON array of job objects:
      {"gid": <int>, "playlist": <str>, "format": "long"|"short",
       "duration": <seconds>, "title": <str>, "channel": <str>,
       "cover": <path or "">, "frame": <path or "">, "video_url": <str>}

  out/_summaries/chunk_*.json  one or more JSON arrays of summary objects,
      keyed by the same gid:
      {"gid": <int>, "about": <str>, "why_it_works": <str>, "recreation_angle": <str>}

Set OUTPUT_DIR to override the default "out" location for both inputs and the
finished workbook.
"""
import glob, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

ROOT = os.environ.get("OUTPUT_DIR", "out")
BOX_W, BOX_H = 230, 130          # display box for thumbnails (px)
ROW_PX = BOX_H + 8

jobs_path = os.path.join(ROOT, "brain_jobs.json")
if not os.path.exists(jobs_path):
    raise SystemExit(f"Missing {jobs_path} -- see this script's docstring for the expected shape.")
jobs = {j["gid"]: j for j in json.load(open(jobs_path))}
summ = {}
for fp in glob.glob(os.path.join(ROOT, "_summaries", "chunk_*.json")):
    try:
        for o in json.load(open(fp)):
            summ[int(o["gid"])] = o
    except Exception as e:
        print("WARN bad summary file", fp, e)
missing = [g for g in jobs if g not in summ]
print(f"videos={len(jobs)} summaries={len(summ)} missing={len(missing)}")
if missing:
    print("MISSING gids:", missing)

def fmt_len(s):
    try:
        s = int(float(s))
    except (ValueError, TypeError):
        return ""
    h, m, sec = s // 3600, (s % 3600) // 60, s % 60
    return f"{h}:{m:02d}:{sec:02d}" if h else f"{m}:{sec:02d}"

COLS = [("thumb", "Thumbnail", 34), ("gid", "#", 5), ("playlist", "Playlist", 20),
        ("format", "Format", 8), ("length", "Length", 8), ("title", "Title", 40),
        ("channel", "Channel", 18), ("about", "About", 46), ("why", "Why it works", 42),
        ("angle", "Recreation angle", 46), ("liked", "Liked?", 8),
        ("what", "What I liked (be specific)", 38), ("recreate", "Recreate?", 11),
        ("twist", "My twist", 30), ("priority", "Priority", 9),
        ("status", "Status", 13), ("link", "Link", 16)]
EDIT = {"liked", "what", "recreate", "twist", "priority", "status"}

wb = Workbook()
ws = wb.active
ws.title = "Vault"
HF = PatternFill("solid", start_color="1F3864")
EDITFILL = PatternFill("solid", start_color="FFF2CC")
HFONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
CELL = Font(name="Arial", size=10)
LINK = Font(name="Arial", size=10, color="0563C1", underline="single")
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.append([c[1] for c in COLS])
for ci, (k, lbl, w) in enumerate(COLS, 1):
    c = ws.cell(1, ci); c.fill = HF; c.font = HFONT
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.row_dimensions[1].height = 28

tmpdir = os.path.join(ROOT, "_thumbtmp"); os.makedirs(tmpdir, exist_ok=True)

def thumb_path(j):
    for p in (j.get("cover"), j.get("frame")):
        if p and os.path.exists(p) and os.path.getsize(p) > 800:
            return p
    return None

r = 1
for gid in sorted(jobs):
    r += 1
    j = jobs[gid]; s = summ.get(gid, {})
    vals = {"gid": gid, "playlist": j["playlist"], "format": j["format"],
            "length": fmt_len(j["duration"]), "title": j["title"], "channel": j["channel"],
            "about": s.get("about", ""), "why": s.get("why_it_works", ""),
            "angle": s.get("recreation_angle", ""), "liked": "", "what": "", "recreate": "",
            "twist": "", "priority": "", "status": "", "link": "watch"}
    for ci, (k, _, _) in enumerate(COLS, 1):
        if k == "thumb":
            continue
        c = ws.cell(r, ci, vals.get(k, ""))
        c.font = CELL
        c.alignment = Alignment(vertical="top", wrap_text=k in ("title", "about", "why", "angle", "what", "twist"))
        c.border = BORDER
        if k in EDIT:
            c.fill = EDITFILL
        if k == "link":
            c.value = "watch"; c.hyperlink = j["video_url"]; c.font = LINK
    # thumbnail image
    tp = thumb_path(j)
    if tp:
        im = PILImage.open(tp).convert("RGB")
        w0, h0 = im.size
        scale = min(BOX_W / w0, BOX_H / h0)
        dw, dh = int(w0 * scale), int(h0 * scale)
        im.thumbnail((max(dw, BOX_W), max(dh, BOX_H)))
        out = os.path.join(tmpdir, f"{gid:03d}.jpg")
        im.save(out, "JPEG", quality=80)
        xim = XLImage(out); xim.width, xim.height = dw, dh
        ws.add_image(xim, f"A{r}")
    else:
        c = ws.cell(r, 1, "(no image)"); c.font = Font(name="Arial", size=9, italic=True, color="999999")
        c.alignment = Alignment(vertical="center", horizontal="center")
    ws.cell(r, 1).border = BORDER
    ws.row_dimensions[r].height = ROW_PX * 0.75

# dropdowns
last = ws.max_row
def col_of(key): return get_column_letter([c[0] for c in COLS].index(key) + 1)
DVS = {"recreate": '"Yes,Maybe,No"', "priority": '"1,2,3"',
       "status": '"Idea,Scripting,Filmed,Published"', "liked": '"y,n"'}
for key, f in DVS.items():
    dv = DataValidation(type="list", formula1=f, allow_blank=True)
    ws.add_data_validation(dv)
    col = col_of(key); dv.add(f"{col}2:{col}{last}")

ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{last}"
out = os.path.join(ROOT, "content_brain_vault.xlsx")
wb.save(out)
print("saved", out, "rows", last - 1)
